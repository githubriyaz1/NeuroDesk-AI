from pathlib import Path
from typing import Any, Dict, List
import openpyxl

from app.core.logging import logger
from app.models.asset import Asset
from app.schemas.preview import PreviewType
from app.services.preview_providers.base import PreviewProvider


class ExcelPreviewProvider(PreviewProvider):
    """Preview provider for Microsoft Excel (.xlsx, .xlsm, .xltx) spreadsheets."""

    @property
    def preview_type(self) -> PreviewType:
        return PreviewType.EXCEL

    def can_handle(self, asset_type: str, mime_type: str, extension: str) -> bool:
        ext = extension.lower().lstrip(".")
        mime = (mime_type or "").lower()
        return ext in ["xlsx", "xls", "xlsm", "xltx"] or "excel" in mime or "spreadsheetml" in mime

    def extract_metadata(self, file_path: Path, asset: Asset) -> Dict[str, Any]:
        try:
            wb = openpyxl.load_workbook(str(file_path), read_only=True, data_only=True)
            sheet_names = wb.sheetnames
            active_sheet = wb.active.title if wb.active else (sheet_names[0] if sheet_names else "Sheet1")
            wb.close()
            return {
                "sheet_names": sheet_names,
                "sheet_count": len(sheet_names),
                "active_sheet": active_sheet,
                "file_size": asset.file_size,
            }
        except Exception as e:
            logger.warning(f"Failed to extract Excel metadata for {file_path}: {e}")
            return {
                "sheet_names": ["Sheet1"],
                "sheet_count": 1,
                "active_sheet": "Sheet1",
                "file_size": asset.file_size,
            }

    def generate_preview(self, file_path: Path, asset: Asset) -> Dict[str, Any]:
        sheets_data: Dict[str, Any] = {}
        sheet_names: List[str] = []
        active_sheet_name = "Sheet1"

        try:
            wb = openpyxl.load_workbook(str(file_path), read_only=True, data_only=True)
            sheet_names = wb.sheetnames
            if sheet_names:
                active_sheet_name = wb.active.title if wb.active else sheet_names[0]

            for sheet_name in sheet_names[:5]:  # Preview up to 5 sheets
                ws = wb[sheet_name]
                rows_data: List[List[Any]] = []
                for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
                    if row_idx >= 50:  # Max 50 rows per sheet
                        break
                    formatted_row = [str(val) if val is not None else "" for val in row]
                    if any(formatted_row):  # Skip completely empty rows
                        rows_data.append(formatted_row)

                headers = rows_data[0] if rows_data else []
                sample_rows = rows_data[1:] if len(rows_data) > 1 else []

                sheets_data[sheet_name] = {
                    "headers": headers,
                    "rows": sample_rows,
                    "total_rows": len(rows_data),
                    "column_count": len(headers),
                }

            wb.close()
        except Exception as e:
            logger.error(f"Failed to generate Excel preview for {file_path}: {e}")
            sheets_data["Sheet1"] = {"headers": [], "rows": [], "total_rows": 0, "column_count": 0}

        return {
            "sheet_names": sheet_names,
            "active_sheet": active_sheet_name,
            "sheets": sheets_data,
        }
